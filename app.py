import os
import io
import csv
import json
import base64
import secrets
from datetime import datetime, date, timedelta
from functools import wraps
import openpyxl
from flask import Flask, request, redirect, url_for, flash, render_template, session, send_file, abort
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import qrcode
from apscheduler.schedulers.background import BackgroundScheduler

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'change-this-secret-before-production-' + secrets.token_hex(8))
db_url = os.getenv('DATABASE_URL', 'sqlite:///kolfe_coupons.db')
if db_url.startswith('postgres://'):
    db_url = db_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'

try:
    from convertdate import ethiopian
except Exception:
    ethiopian = None

class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(100), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default='employee')
    full_name = db.Column(db.String(200), nullable=False, default='')
    employee_id = db.Column(db.String(80), unique=True, index=True)
    sector = db.Column(db.String(150), nullable=False, default='')
    must_change_password = db.Column(db.Boolean, default=True, nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    
    def set_password(self, value):
        self.password_hash = generate_password_hash(value)
        
    def check_password(self, value):
        return check_password_hash(self.password_hash, value)

class Allocation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
    month_key = db.Column(db.String(20), nullable=False, index=True)
    allocated = db.Column(db.Integer, nullable=False, default=0)
    spent = db.Column(db.Integer, nullable=False, default=0)
    expires_at = db.Column(db.Date, nullable=False)
    closed = db.Column(db.Boolean, default=False, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    user = db.relationship('User', backref='allocations')
    __table_args__ = (db.UniqueConstraint('user_id', 'month_key', name='uq_user_month'),)
    
    @property
    def remaining(self):
        return max(0, self.allocated - self.spent)

class SectorBudget(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sector = db.Column(db.String(150), unique=True, nullable=False)
    monthly_amount = db.Column(db.Integer, nullable=False, default=0)

class Redemption(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    employee_user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
    merchant_user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
    allocation_id = db.Column(db.Integer, db.ForeignKey('allocation.id'), nullable=False)
    amount = db.Column(db.Integer, nullable=False)
    reference = db.Column(db.String(100), unique=True, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    employee = db.relationship('User', foreign_keys=[employee_user_id])
    merchant = db.relationship('User', foreign_keys=[merchant_user_id])
    allocation = db.relationship('Allocation')

@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))

def role_required(*roles):
    def deco(fn):
        @wraps(fn)
        @login_required
        def wrapped(*args, **kwargs):
            if current_user.role not in roles:
                abort(403)
            if current_user.must_change_password and request.endpoint not in ('change_password', 'logout', 'static'):
                return redirect(url_for('change_password'))
            return fn(*args, **kwargs)
        return wrapped
    return deco

def ethiopian_info(d=None):
    d = d or date.today()
    if ethiopian:
        ey, em, ed = ethiopian.from_gregorian(d.year, d.month, d.day)
        names = ['መስከረም', 'ጥቅምት', 'ኅዳር', 'ታኅሳስ', 'ጥር', 'የካቲት', 'መጋቢት', 'ሚያዝያ', 'ግንቦት', 'ሰኔ', 'ሐምሌ', 'ነሐሴ', 'ጳጉሜ']
        month_key = f'{ey:04d}-{em:02d}'
        if em == 13:
            ny, nm = ey + 1, 1
        elif em == 12:
            ny, nm = ey, 13
        else:
            ny, nm = ey, em + 1
        next_g = ethiopian.to_gregorian(ny, nm, 1)
        expiry = date(*next_g) - timedelta(days=1)
        return {'year': ey, 'month': em, 'day': ed, 'month_name': names[em-1], 'key': month_key, 'expiry': expiry}
    first_next = (d.replace(day=28) + timedelta(days=4)).replace(day=1)
    return {'year': d.year, 'month': d.month, 'day': d.day, 'month_name': d.strftime('%B'), 'key': d.strftime('%Y-%m'), 'expiry': first_next - timedelta(days=1)}

def qr_data(user):
    return f'Full Name: {user.full_name}\nID: {user.employee_id or ""}\nSector: {user.sector}'

def qr_data_uri(user):
    img = qrcode.make(qr_data(user))
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode()

def current_allocation(user):
    info = ethiopian_info()
    return Allocation.query.filter_by(user_id=user.id, month_key=info['key']).first()

def create_allocation(user, amount, info=None):
    info = info or ethiopian_info()
    existing = Allocation.query.filter_by(user_id=user.id, month_key=info['key']).first()
    if existing:
        existing.allocated = int(amount)
        existing.expires_at = info['expiry']
        existing.closed = False
        return existing
    obj = Allocation(user_id=user.id, month_key=info['key'], allocated=int(amount), spent=0, expires_at=info['expiry'])
    db.session.add(obj)
    return obj

def monthly_rollover():
    with app.app_context():
        info = ethiopian_info()
        Allocation.query.filter(Allocation.month_key != info['key'], Allocation.closed == False).update({'closed': True}, synchronize_session=False)
        budgets = {b.sector: b.monthly_amount for b in SectorBudget.query.all()}
        employees = User.query.filter_by(role='employee', active=True).all()
        for u in employees:
            if not Allocation.query.filter_by(user_id=u.id, month_key=info['key']).first():
                amount = budgets.get(u.sector, 0)
                create_allocation(u, amount, info)
        db.session.commit()

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        u = User.query.filter_by(username=request.form.get('username', '').strip()).first()
        if u and u.active and u.check_password(request.form.get('password', '')):
            login_user(u)
            return redirect(url_for('change_password' if u.must_change_password else 'dashboard'))
        flash('Invalid username/password or inactive account.')
    return render_template('login.html', title='Login')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/change-password', methods=['GET', 'POST'])
@login_required
def change_password():
    if request.method == 'POST':
        old = request.form.get('old_password', '')
        new = request.form.get('new_password', '')
        if not current_user.check_password(old):
            flash('Current password is incorrect.')
        elif len(new) < 8:
            flash('New password must contain at least 8 characters.')
        else:
            current_user.set_password(new)
            current_user.must_change_password = False
            db.session.commit()
            flash('Password changed successfully.')
            return redirect(url_for('dashboard'))
    return render_template('change_password.html', title='Change password')

@app.route('/')
@login_required
def dashboard():
    if current_user.must_change_password:
        return redirect(url_for('change_password'))
    info = ethiopian_info()
    if current_user.role == 'employee':
        a = current_allocation(current_user)
        history = Redemption.query.filter_by(employee_user_id=current_user.id).order_by(Redemption.created_at.desc()).limit(20).all()
        return render_template('employee_dashboard.html', title='My balance', info=info, a=a, history=history)
    if current_user.role in ('merchant', 'authority'):
        total = db.session.query(db.func.coalesce(db.func.sum(Redemption.amount), 0)).filter_by(merchant_user_id=current_user.id).scalar()
        recent = Redemption.query.filter_by(merchant_user_id=current_user.id).order_by(Redemption.created_at.desc()).limit(20).all()
        return render_template('operator_dashboard.html', title='Operator dashboard', total=total, recent=recent)
    counts = {r: User.query.filter_by(role=r, active=True).count() for r in ('employee', 'merchant', 'authority', 'admin')}
    info = ethiopian_info()
    spent = db.session.query(db.func.coalesce(db.func.sum(Allocation.spent), 0)).filter_by(month_key=info['key']).scalar()
    alloc = db.session.query(db.func.coalesce(db.func.sum(Allocation.allocated), 0)).filter_by(month_key=info['key']).scalar()
    return render_template('admin_dashboard.html', title='Admin dashboard', counts=counts, info=info, spent=spent, alloc=alloc)

@app.route('/setup', methods=['GET', 'POST'])
def setup_system():
    if User.query.filter_by(role='admin').first():
        flash('Admin account already exists! Please log in.')
        return redirect(url_for('login'))
    if request.method == 'POST':
        username = request.form.get('username', 'admin').strip()
        password = request.form.get('password', '').strip()
        if len(password) < 8:
            flash('Password must be at least 8 characters.')
        else:
            u = User(username=username, full_name='System Administrator', role='admin', must_change_password=False)
            u.set_password(password)
            db.session.add(u)
            db.session.commit()
            flash('Admin account created successfully! Please log in.')
            return redirect(url_for('login'))
    return '''
        <!doctype html>
        <title>System Setup</title>
        <body style="font-family:Arial; padding:40px; max-width:400px; margin:auto; background:#f3f6f9;">
            <div style="background:white; padding:20px; border-radius:8px; box-shadow:0 2px 10px rgba(0,0,0,0.1);">
                <h2>System Setup</h2>
                <p>Create your initial administrator account:</p>
                <form method="post">
                    <label style="font-weight:bold; display:block; margin-top:10px;">Admin Username:</label>
                    <input type="text" name="username" value="admin" style="width:100%; padding:8px; margin-top:5px; border:1px solid #ccc; border-radius:4px;">
                    <label style="font-weight:bold; display:block; margin-top:10px;">Admin Password:</label>
                    <input type="password" name="password" placeholder="At least 8 characters" style="width:100%; padding:8px; margin-top:5px; border:1px solid #ccc; border-radius:4px;">
                    <button type="submit" style="width:100%; padding:10px; background:#176b87; color:#fff; border:none; border-radius:4px; font-weight:bold; margin-top:15px; cursor:pointer;">Create Admin Account</button>
                </form>
            </div>
        </body>
    '''

@app.route('/admin/employees')
@role_required('admin')
def employees():
    employees = User.query.filter_by(role='employee', active=True).order_by(User.employee_id).all()
    return render_template('employees.html', title='Employees', employees=employees)

@app.route('/admin/upload_employees', methods=['GET', 'POST'])
@role_required('admin')
def upload_employees():
    if request.method == 'GET':
        return redirect(url_for('employees'))
        
    file = request.files.get('file')
    if not file or not file.filename:
        flash('ፋይል አልተመረጠም', 'danger')
        return redirect(url_for('employees'))

    try:
        filename_lower = file.filename.lower()
        added_count = 0
        rows = []
        file_bytes = file.read()

        if filename_lower.endswith('.csv'):
            decoded_text = None
            for enc in ['utf-8-sig', 'utf-8', 'cp1252', 'latin-1']:
                try:
                    decoded_text = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if not decoded_text:
                decoded_text = file_bytes.decode('utf-8', errors='ignore')
            
            stream = io.StringIO(decoded_text)
            reader = list(csv.reader(stream))
            if not reader:
                flash('ባዶ ፋይል ነው', 'danger')
                return redirect(url_for('employees'))
            
            header_row = [str(h).strip().lower() for h in reader[0]]
            id_idx = -1
            name_idx = -1
            sector_idx = -1
            
            for i, h in enumerate(header_row):
                if ('id' in h) and id_idx == -1: id_idx = i
                elif ('name' in h or 'ስም' in h) and name_idx == -1: name_idx = i
                elif ('sector' in h or 'department' in h or 'መስሪያ' in h or 'ክፍል' in h) and sector_idx == -1: sector_idx = i

            start_row_idx = 1
            if id_idx == -1 or name_idx == -1:
                id_idx = 0
                name_idx = 1
                sector_idx = 2 if len(header_row) > 2 else -1
                start_row_idx = 0

            for row in reader[start_row_idx:]:
                if row and len(row) > max(id_idx, name_idx):
                    emp_id = row[id_idx].strip()
                    full_name = row[name_idx].strip()
                    sector = row[sector_idx].strip() if sector_idx != -1 and len(row) > sector_idx and row[sector_idx].strip() else 'ከተማ ውበት'
                    if emp_id and full_name and emp_id.lower() != 'nan':
                        rows.append((emp_id, full_name, sector))

        elif filename_lower.endswith(('.xlsx', '.xls')):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet = wb.active
            iter_rows = list(sheet.iter_rows(values_only=True))
            if not iter_rows:
                flash('ባዶ ፋይል ነው', 'danger')
                return redirect(url_for('employees'))

            header_row = [str(h).strip().lower() if h is not None else '' for h in iter_rows[0]]
            id_idx = -1
            name_idx = -1
            sector_idx = -1

            for i, h in enumerate(header_row):
                if ('id' in h) and id_idx == -1: id_idx = i
                elif ('name' in h or 'ስም' in h) and name_idx == -1: name_idx = i
                elif ('sector' in h or 'department' in h or 'መስሪያ' in h or 'ክፍል' in h) and sector_idx == -1: sector_idx = i

            start_row_idx = 1
            if id_idx == -1 or name_idx == -1:
                id_idx = 0
                name_idx = 1
                sector_idx = 2 if len(header_row) > 2 else -1
                start_row_idx = 0

            for row in iter_rows[start_row_idx:]:
                if row and len(row) > max(id_idx, name_idx):
                    val_id = row[id_idx]
                    val_name = row[name_idx]
                    val_sector = row[sector_idx] if sector_idx != -1 and len(row) > sector_idx else 'ከተማ ውበት'
                    
                    emp_id = str(val_id).strip() if val_id is not None else ''
                    full_name = str(val_name).strip() if val_name is not None else ''
                    sector = str(val_sector).strip() if val_sector is not None else 'ከተማ ውበት'

                    if emp_id and full_name and emp_id.lower() != 'nan':
                        rows.append((emp_id, full_name, sector))
            wb.close()
        else:
            flash('እባክዎ ትክክለኛ የ CSV ወይም Excel (.xlsx) ፋይል ይምረጡ', 'danger')
            return redirect(url_for('employees'))

        for emp_id, full_name, sector in rows:
            existing = User.query.filter_by(employee_id=emp_id).first()
            if not existing:
                username = f'emp_{emp_id}'
                new_user = User(
                    username=username,
                    employee_id=emp_id,
                    full_name=full_name,
                    sector=sector,
                    role='employee',
                    must_change_password=True
                )
                new_user.set_password('123456')
                db.session.add(new_user)
                added_count += 1

        db.session.commit()
        flash(f'በተሳካ ሁኔታ {added_count} ሰራተኞች ተጭነዋል!', 'success')
    except Exception as e:
        flash(f'ፋይልን በማንበብ ጊዜ ስህተት ተፈጥሯል: {str(e)}', 'danger')

    return redirect(url_for('employees'))

@app.route('/health')
def health():
    return {'status': 'ok'}

with app.app_context():
    db.create_all()
    admin_username = os.getenv('ADMIN_USERNAME', 'admin')
    admin_password = os.getenv('ADMIN_PASSWORD', 'ChangeMe123!')
    if not User.query.filter_by(role='admin').first():
        u = User(username=admin_username, full_name='System Administrator', role='admin', must_change_password=True)
        u.set_password(admin_password or secrets.token_urlsafe(12))
        db.session.add(u)
        db.session.commit()

scheduler = BackgroundScheduler(daemon=True)
scheduler.add_job(monthly_rollover, 'cron', hour=0, minute=10, id='ethiopian_month_rollover', replace_existing=True)
scheduler.start()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', 5000)), debug=False)